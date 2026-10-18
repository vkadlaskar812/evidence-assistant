import csv, io, json, os, re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, Form, HTTPException
from fastapi.responses import PlainTextResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from validator import validate_claim
import time

BASE = Path(__file__).parent
load_dotenv(BASE / ".env", override=True)
EMP = os.getenv("EMPLOYEE", "Priya")
app = FastAPI()
STATE = {"artifacts": json.loads((BASE / os.getenv("DATA_FILE", "data/demo_data.json")).read_text(encoding="utf-8")), "claims": {}, "answers": {}}

PROMPT = f"""You prepare the annual promotion case for an employee named {EMP}. Return a JSON list.
Each item: project, work_done (1-2 sentences on what {EMP} personally did), impact (the OUTCOME of the work,
only if a source states it, otherwise an empty string; never restate the activity), artifact_ids (ids you used).
Rules:
- Never invent numbers or facts.
- Every source has an owner field. Credit {EMP} only for sources whose owner is {EMP}. Work owned by anyone else,
  or by "Team", is not hers: skip it, or describe it as a team result beginning "Team result:".
- Team-level or epic-level figures are not her own impact: leave impact empty or begin it with "Team result:".
- Skip open or unfinished items and trivial changes such as typos.
- Counts such as attendees, tests, pages or tickets describe the activity; they are not impact. Leave impact empty unless a source states a result (time saved, usage, quality, revenue).
- Combine only closely related records (a ticket with its pull request and design note) into one item per deliverable. Never bundle unrelated work into one item. Aim for 3 to 6 items for the evidence below.
Evidence:
"""

CHECK = """You are a fact-checker. Using ONLY the sources (title and excerpt both count as source text), decide whether the claim adds anything the sources do not support. Flag only substantive additions: new facts or numbers, causal links ("by doing X", "which caused"), ownership ("authored", "led", "delivered") that does not match the source's owner field, and team or epic results credited to one person. Do NOT flag paraphrase or details that appear in a title. Return JSON: {"supported": true or false, "issue": "one short sentence, empty if supported"}.
"""

def persist():
    (BASE / "last_state.json").write_text(json.dumps(STATE), encoding="utf-8")
    if os.getenv("USE_FIRESTORE") == "1":
        from google.cloud import firestore
        firestore.Client().collection("cases").document("demo").set({"state": json.dumps(STATE)})

def _ask(prompt):
    from google import genai
    key = os.getenv("GEMINI_API_KEY")
    client = genai.Client(api_key=key) if key else genai.Client(
        vertexai=True, project=os.getenv("GOOGLE_CLOUD_PROJECT"),
        location=os.getenv("GOOGLE_CLOUD_LOCATION", "us-central1"))
    r = client.models.generate_content(model=os.getenv("GEMINI_MODEL", "gemini-2.5-flash"),
        contents=prompt, config={"response_mime_type": "application/json", "temperature": 0})
    return json.loads(r.text)

def gemini_claims(arts):
    groups = {}
    for a in arts: groups.setdefault(a["project"], []).append(a)
    failed = []
    def one(item):
        proj, items = item
        for attempt in range(3):
            try: return _ask(PROMPT + json.dumps(items))
            except Exception as e:
                print(f"extraction failed for {proj} (attempt {attempt + 1}): {e}")
                time.sleep(4 * (attempt + 1))
        failed.append(proj); return []
    with ThreadPoolExecutor(max_workers=2) as ex:
        res = [c for r in ex.map(one, groups.items()) for c in r]
    return res, failed

def wording_check(claim, by_id):
    srcs = [{k: by_id[i].get(k, "") for k in ("id", "owner", "title", "excerpt")} for i in claim["artifact_ids"] if i in by_id]
    try:
        return _ask(CHECK + json.dumps({"claim": claim["work_done"] + " " + claim["impact"], "sources": srcs}))
    except Exception:
        return None  # if the check fails, change nothing

def fallback_claims(arts):  # no-key mode so the UI can be tested; NOT the real extraction
    out = {}
    for a in arts:
        c = out.setdefault(a["project"], {"project": a["project"], "work_done": "", "impact": "", "artifact_ids": []})
        c["work_done"] += a["title"] + ". "; c["artifact_ids"].append(a["id"])
    return list(out.values())

def revalidate():
    by_id = {a["id"]: a for a in STATE["artifacts"]}
    for cid, c in STATE["claims"].items():
        c["status"], c["note"] = validate_claim(c, by_id, STATE["answers"].get(cid))
        w = c.get("wording")
        if w and w.get("supported") is False and c["status"] != "unsupported":
            c["status"], c["note"] = "unsupported", "wording goes beyond sources: " + w.get("issue", "")
        elif w and w.get("supported") is True and c["status"] == "unchecked":
            c["status"], c["note"] = "verified", "no figures; wording checked by AI"
        c["question"] = "" if c["impact"] else f"What was the measurable effect of this work on {c['project']}?"

@app.get("/api/state")
def state(): return STATE

@app.post("/api/extract")
def extract():
    live = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_CLOUD_PROJECT")
    failed = []
    if live: raw, failed = gemini_claims(STATE["artifacts"])
    else: raw = fallback_claims(STATE["artifacts"])
    STATE["claims"] = {f"C{i+1}": {**c, "id": f"C{i+1}", "decision": "pending"} for i, c in enumerate(raw)}
    STATE["answers"] = {}
    revalidate()
    if live:
        by_id = {a["id"]: a for a in STATE["artifacts"]}
        claims = list(STATE["claims"].values())
        with ThreadPoolExecutor(max_workers=4) as ex:
            for c, w in zip(claims, ex.map(lambda c: wording_check(c, by_id), claims)):
                c["wording"] = w
        revalidate()
    persist()
    return {"mode": "gemini" if live else "fallback", "claims": len(raw), "failed": failed}

class Answer(BaseModel): text: str
@app.post("/api/claims/{cid}/answer")
def answer(cid: str, a: Answer):
    if cid not in STATE["claims"]: raise HTTPException(404)
    STATE["answers"][cid] = a.text; STATE["claims"][cid]["impact"] = a.text
    revalidate(); persist(); return STATE["claims"][cid]

class Decision(BaseModel): decision: str
@app.post("/api/claims/{cid}/decision")
def decide(cid: str, d: Decision):
    if cid not in STATE["claims"]: raise HTTPException(404)
    STATE["claims"][cid]["decision"] = d.decision; persist(); return STATE["claims"][cid]

@app.post("/api/upload")
async def upload(file: UploadFile, project: str = Form("Uploaded")):
    data, name = await file.read(), file.filename.lower()
    if name.endswith(".docx"):
        from docx import Document
        text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    elif name.endswith(".pptx"):
        from pptx import Presentation
        text = "\n".join(s.text_frame.text for sl in Presentation(io.BytesIO(data)).slides
                         for s in sl.shapes if s.has_text_frame)
    else: text = data.decode("utf-8", "ignore")
    aid = f"UP-{len(STATE['artifacts']) + 1}"
    STATE["artifacts"].append({"id": aid, "source_type": "upload", "owner": EMP, "project": project, "date": "",
        "title": file.filename, "url": "", "excerpt": re.sub(r"\s+", " ", text)[:2000]})
    return {"id": aid}

@app.get("/api/export.csv", response_class=PlainTextResponse)
def export():
    buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(["Project", "Work done", "Impact", "Verification", "Evidence"])
    by_id = {a["id"]: a for a in STATE["artifacts"]}
    for c in STATE["claims"].values():
        if c["decision"] == "approve" and c["status"] != "unsupported":
            w.writerow([c["project"], c["work_done"], c["impact"], c["status"],
                        "; ".join(by_id[i]["url"] or i for i in c["artifact_ids"] if i in by_id)])
    return buf.getvalue()

app.mount("/", StaticFiles(directory=BASE / "static", html=True), name="static")